from datetime import date, timedelta
import sqlite3
import pytest
from openpyxl import load_workbook
from app.services.backup import BackupService
from app.services.reports import ReportService


def test_backup_restore_and_integrity(finance, tmp_path):
    aid = finance.save_account(finance.sample)
    service = BackupService(finance.db)
    backup = service.create(tmp_path / "backups")
    finance.delete_account(aid)
    assert finance.accounts() == []
    safety = service.restore(backup)
    assert safety.is_file()
    assert finance.account(aid)["nome"] == "Energia elétrica"
    finance.save_account(finance.sample | {"nome": "Após restaurar"})
    assert len(finance.accounts()) == 2
    invalid = tmp_path / "bad.sqlite3"
    invalid.write_bytes(b"invalid backup")
    with pytest.raises((ValueError, sqlite3.DatabaseError)):
        service.restore(invalid)
    assert len(finance.accounts()) == 2


def test_automatic_backup_schedule(finance, tmp_path):
    service = BackupService(finance.db)
    finance.db.set_setting("backup_frequencia", "Diário")
    finance.db.set_setting("backup_pasta", tmp_path / "auto")
    assert service.automatic().is_file()
    assert service.automatic() is None
    finance.db.set_setting(
        "ultimo_backup", (date.today() - timedelta(days=6)).isoformat()
    )
    finance.db.set_setting("backup_frequencia", "Semanal")
    assert service.automatic() is None
    finance.db.set_setting(
        "ultimo_backup", (date.today() - timedelta(days=7)).isoformat()
    )
    assert service.automatic().is_file()


def test_reports_export_real_data(finance, tmp_path):
    aid = finance.save_account(
        finance.sample
        | {"nome": '=HYPERLINK("unsafe")', "data_vencimento": date.today()}
    )
    finance.pay(
        aid,
        {
            "data_pagamento": date.today(),
            "valor_centavos": 125000,
            "forma_pagamento": "PIX",
        },
    )
    reports = ReportService(finance)
    start = date.today() - timedelta(days=31)
    end = date.today() + timedelta(days=31)
    data = reports.report(start, end)
    assert data["Resumo"][1][1] == "R$ 1.250,00"
    xlsx = tmp_path / "report.xlsx"
    pdf = tmp_path / "report.pdf"
    reports.excel(xlsx, start, end)
    reports.pdf(pdf, start, end)
    wb = load_workbook(xlsx)
    assert len(wb.sheetnames) == 7
    assert wb["Pagamentos"]["C3"].value == 1250
    assert wb["Contas"]["B3"].data_type != "f"
    assert pdf.read_bytes().startswith(b"%PDF")
    assert pdf.stat().st_size > 1000
    finance.undo_payment(aid)
    assert reports.report(start, end)["Resumo"][1][1] == "R$ 0,00"
